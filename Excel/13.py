from  clearConsole import clearConsole
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.chart import BarChart, PieChart, Series, Reference

clearConsole()

wb = load_workbook("crime_report.xlsx")
ws = wb.active

chart =  BarChart()
data =  Reference(ws , min_row = 10, min_col = 1 , max_row = 13, max_col =13)
labels = Reference(ws, min_row = 8, min_col = 2 , max_row = 8, max_col =13)
chart.add_data(data, titles_from_data = True)
chart.set_categories(labels)
chart.title = 'Counterfeit Crmies by District'

chart.height = 4.56
chart.width =20.3
ws.add_chart(chart, "B14" )


chart2 = PieChart()
data =  Reference(ws,min_row = 15, min_col = 16 , max_row = 8, max_col =8)
labels = Reference(ws,min_row = 15, min_col = 16, max_row = 7, max_col =7)
chart.add_data(data, from_rows = True)
chart2.set_categories(labels)
chart2.title = ' Counterfeit Crmies by District'

chart2.height = 4.56
chart2.width =8.45
ws.add_chart(chart2, "N14" )

wb.save("lines.xlsx")



print("DONE")