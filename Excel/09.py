# image
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image
from clearConsole import *
from openpyxl import load_workbook

clearConsole()

wb = load_workbook("Pie.xlsx")
ws = wb.active

tab = Table(displayName="Table1", ref="A1:B5")
style= TableStyleInfo(name = "TableStyleMedium9", showFirstColumn= False, showLastColumn= False , 
                        showRowStripes=True, showColumnStripes=True)
            

tab.tableStyleInfo = style
ws.add_table(tab)
wb.save("table.xlsx")

img = Image('kubus.png')
img.height = img.height * .9
img.width = img.width * .9
ws.add_image(img,"A16")
wb.save("Image.xlsx")



print("DONE")

